"""Render tailored résumé content as Markdown, DOCX and tailoring notes."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from .ai import describe_ai


def _joined(*parts: Any) -> str:
    return " | ".join(str(part).strip() for part in parts if str(part or "").strip())


def contact_lines(identity: dict[str, Any]) -> list[str]:
    """The contact lines under the name; empty fields are left out."""
    lines = [
        _joined(identity.get("location"), identity.get("phone"), identity.get("email")),
        _joined(identity.get("linkedin"), identity.get("github"), identity.get("website")),
    ]
    return [line for line in lines if line]


def experience_heading(experience: dict[str, Any]) -> str:
    return _joined(experience.get("role"), experience.get("company"), experience.get("dates"))


def skill_line(skill: dict[str, Any]) -> tuple[str, str]:
    """(category, text); the category may be empty."""
    return str(skill.get("category") or "").strip(), str(skill.get("text") or "").strip()


def markdown_resume(content: dict[str, Any]) -> str:
    identity = content["identity"]
    lines = [f"# {identity['name']}"]
    if identity.get("headline"):
        lines.append(f"**{identity['headline']}**")
    lines += ["", *contact_lines(identity), "", "## PROFESSIONAL SUMMARY", content["summary"], ""]

    if content.get("technical_skills"):
        lines.append("## SKILLS")
        for skill in content["technical_skills"]:
            category, text = skill_line(skill)
            lines.append(f"**{category}:** {text}" if category else f"- {text}")
        lines.append("")

    if content.get("experience"):
        lines.append("## PROFESSIONAL EXPERIENCE")
        for experience in content["experience"]:
            lines.append(f"**{experience_heading(experience)}**")
            for bullet in experience.get("bullets", []):
                lines.append(f"- {bullet['text']}")
            lines.append("")

    if content.get("education"):
        lines.append("## EDUCATION")
        for education in content["education"]:
            lines.append(f"**{education['program']}**")
            if where := _joined(education.get("school"), education.get("dates")):
                lines.append(where)
            for bullet in education.get("bullets", []):
                lines.append(f"- {bullet['text']}")
            lines.append("")

    if content.get("certifications"):
        lines.append("## CERTIFICATIONS")
        primary = content["certifications"][:4]
        additional = content["certifications"][4:]
        for cert in primary:
            lines.append(f"- {cert['text']}")
        if additional:
            lines.append("Additional certificates: " + "; ".join(c["text"] for c in additional))
        lines.append("")

    return "\n".join(lines).rstrip() + "\n"


def _set_cellless_doc_defaults(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0

    if "List Bullet" in styles:
        bullet = styles["List Bullet"]
        bullet.font.name = "Arial"
        bullet.font.size = Pt(9.4)
        bullet.paragraph_format.space_after = Pt(1)
        bullet.paragraph_format.left_indent = Inches(0.18)
        bullet.paragraph_format.first_line_indent = Inches(-0.12)


def _section_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10.5)


def _bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def docx_resume(content: dict[str, Any], path: Path) -> None:
    document = Document()
    _set_cellless_doc_defaults(document)
    identity = content["identity"]

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(identity["name"])
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(16)

    if identity.get("headline"):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(identity["headline"])
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(10.5)

    for line in contact_lines(identity):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        r.font.name = "Arial"
        r.font.size = Pt(8.8)

    _section_heading(document, "PROFESSIONAL SUMMARY")
    document.add_paragraph(content["summary"])

    if content.get("technical_skills"):
        _section_heading(document, "SKILLS")
        for skill in content["technical_skills"]:
            category, text = skill_line(skill)
            if not category:
                _bullet(document, text)
                continue
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(f"{category}: ")
            r.bold = True
            p.add_run(text)

    if content.get("experience"):
        _section_heading(document, "PROFESSIONAL EXPERIENCE")
        for experience in content["experience"]:
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(experience_heading(experience))
            r.bold = True
            for bullet in experience.get("bullets", []):
                _bullet(document, bullet["text"])

    if content.get("education"):
        _section_heading(document, "EDUCATION")
        for education in content["education"]:
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(education["program"])
            r.bold = True
            if where := _joined(education.get("school"), education.get("dates")):
                document.add_paragraph(where)
            for bullet in education.get("bullets", []):
                _bullet(document, bullet["text"])

    if content.get("certifications"):
        _section_heading(document, "CERTIFICATIONS")
        primary = content["certifications"][:4]
        additional = content["certifications"][4:]
        for cert in primary:
            _bullet(document, cert["text"])
        if additional:
            document.add_paragraph("Additional certificates: " + "; ".join(c["text"] for c in additional))

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def tailoring_notes(
    job: dict[str, Any],
    scored: dict[str, Any],
    master_version: str,
    ai_audit: dict[str, Any],
) -> str:
    def values(key: str) -> str:
        items = scored.get(key, [])
        return ", ".join(items) if items else "None"

    return "\n".join(
        [
            "# Tailoring Notes",
            "",
            f"- Master resume version: {master_version}",
            f"- Job: {job.get('title', '')} — {job.get('company', '')}",
            f"- Match score at generation: {scored.get('score', 0)}%",
            f"- Emphasized matched skills: {values('matched_skills')}",
            f"- Missing core skills NOT added to resume: {values('missing_skills')}",
            f"- Optional gaps NOT added to resume: {values('optional_missing_skills')}",
            f"- AI wording: {describe_ai(ai_audit.get('status', ''), ai_audit.get('model', ''))}",
            "",
            "Policy: deterministic selection happens first. AI may improve wording only inside the "
            "master-resume factual whitelist. Rule checks and a second AI factual-verification pass "
            "must succeed before AI wording is accepted; otherwise the deterministic resume is used.",
            "",
        ]
    )
